"""PDF and DOCX exports of reports and revised documents."""

import io
from datetime import datetime
from xml.sax.saxutils import escape

from docx import Document as DocxDocument
from docx.enum.text import WD_COLOR_INDEX
from docx.shared import Pt, RGBColor
from reportlab.lib import colors
from reportlab.lib.enums import TA_LEFT
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import KeepTogether, PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

CATEGORY_COLORS = {
    "unattributed": "#fecaca",
    "cited_no_quotes": "#fde68a",
    "quoted_uncited": "#fde68a",
    "quoted_cited": "#e5e7eb",
    "bibliography": "#e5e7eb",
    "paraphrase": "#fed7aa",
    "paraphrase_cited": "#ffedd5",
}
DOCX_HIGHLIGHT = {
    "unattributed": WD_COLOR_INDEX.RED,
    "cited_no_quotes": WD_COLOR_INDEX.YELLOW,
    "quoted_uncited": WD_COLOR_INDEX.YELLOW,
    "paraphrase": WD_COLOR_INDEX.TAN if hasattr(WD_COLOR_INDEX, "TAN") else WD_COLOR_INDEX.YELLOW,
    "paraphrase_cited": WD_COLOR_INDEX.GRAY_25,
    "quoted_cited": WD_COLOR_INDEX.GRAY_25,
    "bibliography": WD_COLOR_INDEX.GRAY_25,
}


def _styles():
    ss = getSampleStyleSheet()
    return {
        "title": ParagraphStyle("t", parent=ss["Title"], fontSize=20, leading=24, alignment=TA_LEFT, spaceAfter=6),
        "h2": ParagraphStyle("h2", parent=ss["Heading2"], fontSize=13, leading=16, spaceBefore=10, spaceAfter=4, textColor=colors.HexColor("#1e293b")),
        "body": ParagraphStyle("b", parent=ss["BodyText"], fontSize=9.5, leading=13),
        "small": ParagraphStyle("s", parent=ss["BodyText"], fontSize=8, leading=10.5, textColor=colors.HexColor("#475569")),
        "doc": ParagraphStyle("d", parent=ss["BodyText"], fontName="Times-Roman", fontSize=10.5, leading=15),
        "quote": ParagraphStyle("q", parent=ss["BodyText"], fontSize=8.5, leading=11.5, leftIndent=8, textColor=colors.HexColor("#334155"), backColor=colors.HexColor("#f8fafc")),
    }


def _p(text: str, style) -> Paragraph:
    return Paragraph(escape(text or "").replace("\n", "<br/>"), style)


def _highlighted_paragraphs(text: str, findings: list[dict], source_rank: dict, style) -> list:
    """The document text with similarity matches highlighted and numbered."""
    marks = sorted((f for f in findings if f.get("kind") == "similarity"), key=lambda f: f["start"])
    out_html, pos = [], 0
    for f in marks:
        if f["start"] < pos:
            continue
        out_html.append(escape(text[pos : f["start"]]))
        color = CATEGORY_COLORS.get(f["category"], "#e5e7eb")
        rank = source_rank.get(f["source_id"], "?")
        out_html.append(f'<font backColor="{color}">{escape(text[f["start"]:f["end"]])}</font><super><font size="7">[{rank}]</font></super>')
        pos = f["end"]
    out_html.append(escape(text[pos:]))
    html = "".join(out_html)
    return [Paragraph(chunk.replace("\n", "<br/>"), style) for chunk in html.split("\n\n") if chunk.strip()]


def report_pdf(report: dict) -> bytes:
    st = _styles()
    buf = io.BytesIO()
    doc_info = report["document"]
    pdf = SimpleDocTemplate(buf, pagesize=A4, leftMargin=18 * mm, rightMargin=18 * mm, topMargin=16 * mm, bottomMargin=16 * mm, title=f"Integrity report – {doc_info['title']}")
    s = report["summary"]
    sim = report["similarity"]
    story: list = [
        _p("Academic Integrity & Writing Report", st["title"]),
        _p(f"{doc_info['title']}" + (f" — {doc_info['author']}" if doc_info.get("author") else ""), st["body"]),
        _p(f"{doc_info['words']} words · language: {doc_info.get('language_name') or 'unknown'} · generated {report['generated_at'][:16].replace('T', ' ')} UTC", st["small"]),
        Spacer(1, 6),
    ]
    rows = [
        ["Similarity index", f"{s['similarity_index']}%", "Words in exact/near-exact matches needing attention ÷ countable words"],
        ["Unattributed overlap", f"{s['unattributed_pct']}%", "Matches with no quotation marks and no citation"],
        ["Possible paraphrase", f"{s['semantic_pct']}%", "Sentences semantically close to a source (reported separately)"],
        ["Quoted and cited", f"{s['quoted_pct']}%", "Shown for information, not counted"],
        ["AI-writing indicators", s["ai_label"], "Banded result; not proof of authorship"],
        ["Writing quality", f"{s['quality_score']}/100", "Average of four disclosed subscores"],
        ["Citation issues", str(s["citation_issues"]), "Missing references, uncited quotes, unverifiable DOIs"],
    ]
    t = Table([[_p(a, st["body"]), _p(b, st["body"]), _p(c, st["small"])] for a, b, c in rows], colWidths=[45 * mm, 30 * mm, 99 * mm])
    t.setStyle(TableStyle([("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#cbd5e1")), ("VALIGN", (0, 0), (-1, -1), "TOP"), ("BACKGROUND", (0, 0), (0, -1), colors.HexColor("#f1f5f9"))]))
    story += [t, Spacer(1, 6), _p("Similarity is not proof of plagiarism. Every match should be reviewed in context.", st["small"])]

    cov = report["coverage"]
    story += [_p("What was searched", st["h2"]), _p(cov.get("summary", ""), st["body"])]
    prov_rows = [[_p("Service", st["small"]), _p("Status", st["small"]), _p("Requests", st["small"]), _p("Results", st["small"])]]
    for p in cov.get("providers", []):
        prov_rows.append([_p(p["label"], st["small"]), _p(p["status"] + (f" – {p.get('detail')}" if p.get("detail") else ""), st["small"]), _p(str(p.get("requests", "–")), st["small"]), _p(str(p.get("results", "–")), st["small"])])
    pt = Table(prov_rows, colWidths=[62 * mm, 72 * mm, 20 * mm, 20 * mm])
    pt.setStyle(TableStyle([("GRID", (0, 0), (-1, -1), 0.3, colors.HexColor("#cbd5e1")), ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#f1f5f9"))]))
    story += [pt, Spacer(1, 3)]
    for line in cov.get("not_covered", []):
        story.append(_p("Not covered: " + line, st["small"]))

    story.append(_p("Matched sources", st["h2"]))
    if sim["sources"]:
        src_rows = [[_p("#", st["small"]), _p("Source", st["small"]), _p("Counted", st["small"]), _p("Paraphrase", st["small"])]]
        for r in sim["sources"]:
            where = r.get("doi") and f"doi:{r['doi']}" or r.get("url") or ""
            src_rows.append([_p(str(r["rank"]), st["small"]), _p(f"{r['title']}\n{where}\n({r['provider']}, {r['text_kind'].replace('_', ' ')})", st["small"]), _p(f"{r['percent']}%", st["small"]), _p(f"{r['semantic_percent']}%", st["small"])])
        stbl = Table(src_rows, colWidths=[8 * mm, 122 * mm, 22 * mm, 22 * mm])
        stbl.setStyle(TableStyle([("GRID", (0, 0), (-1, -1), 0.3, colors.HexColor("#cbd5e1")), ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#f1f5f9")), ("VALIGN", (0, 0), (-1, -1), "TOP")]))
        story.append(stbl)
    else:
        story.append(_p("No matching sources were found among the documents that could be searched.", st["body"]))

    rank = {r["id"]: r["rank"] for r in sim["sources"]}
    story.append(_p("Findings", st["h2"]))
    for f in sim["findings"]:
        block = [
            _p(f"[{rank.get(f['source_id'], '?')}] {f['label']} · {f['words']} words · {f['method'].replace('_', ' ')} · severity {f['severity']}", st["body"]),
            _p(report["text"][f["start"] : f["end"]][:1200], st["quote"]),
        ]
        if f.get("source_excerpt"):
            block.append(_p("Source: " + f["source_excerpt"][:1200], st["quote"]))
        block.append(_p(f["explanation"] + " " + f["next_steps"], st["small"]))
        block.append(Spacer(1, 5))
        story.append(KeepTogether(block))
    if not sim["findings"]:
        story.append(_p("No similarity findings.", st["body"]))

    ai = report["ai"]
    story += [_p("AI-writing indicators", st["h2"]), _p(f"Result: {ai['label']}. {ai.get('explanation', '')}", st["body"])]
    for f in ai.get("features", [])[:8]:
        story.append(_p(f"{f['label']}: {f['value']} (typical human {f.get('typical_human')}, typical AI {f.get('typical_ai')}; contribution {f['contribution']:+.2f})", st["small"]))
    for c in ai.get("caveats", []):
        story.append(_p("• " + c, st["small"]))

    q = report["quality"]
    story += [_p("Writing quality", st["h2"]), _p(f"Quality index {q['score']}/100 — " + ", ".join(f"{k} {v}" for k, v in q["subscores"].items()), st["body"])]
    rd = q["readability"]
    story.append(_p(f"Flesch reading ease {rd['flesch_reading_ease']}, grade level {rd['flesch_kincaid_grade']}, average sentence {rd['avg_sentence_words']} words.", st["small"]))
    for k, v in q["formulas"].items():
        story.append(_p(f"{k}: {v}", st["small"]))

    c = report["citations"]
    story += [_p("Citations and references", st["h2"]), _p(f"Style: {c.get('style') or 'none detected'} · {len(c['citations'])} in-text citations · {len(c['references'])} references.", st["body"])]
    for issue in c["issues"][:60]:
        story.append(_p(f"• {issue['message']}", st["small"]))

    story += [PageBreak(), _p("Document with highlighted matches", st["h2"])]
    story += _highlighted_paragraphs(report["text"], sim["findings"], rank, st["doc"])
    pdf.build(story)
    return buf.getvalue()


def report_docx(report: dict) -> bytes:
    d = DocxDocument()
    d.styles["Normal"].font.size = Pt(10.5)
    info, s = report["document"], report["summary"]
    d.add_heading("Academic Integrity & Writing Report", 0)
    d.add_paragraph(f"{info['title']}" + (f" — {info['author']}" if info.get("author") else ""))
    d.add_paragraph(f"{info['words']} words · generated {report['generated_at'][:16].replace('T', ' ')} UTC")
    table = d.add_table(rows=0, cols=2)
    table.style = "Light Grid Accent 1"
    for k, v in [
        ("Similarity index", f"{s['similarity_index']}%"),
        ("Unattributed overlap", f"{s['unattributed_pct']}%"),
        ("Possible paraphrase", f"{s['semantic_pct']}%"),
        ("Quoted and cited", f"{s['quoted_pct']}%"),
        ("AI-writing indicators", s["ai_label"]),
        ("Writing quality", f"{s['quality_score']}/100"),
        ("Citation issues", str(s["citation_issues"])),
    ]:
        row = table.add_row().cells
        row[0].text, row[1].text = k, v
    d.add_paragraph("Similarity is not proof of plagiarism; review each match in context.")
    d.add_heading("What was searched", 1)
    d.add_paragraph(report["coverage"].get("summary", ""))
    for p in report["coverage"].get("providers", []):
        d.add_paragraph(f"{p['label']}: {p['status']}" + (f" ({p.get('detail')})" if p.get("detail") else ""), style="List Bullet")
    d.add_heading("Matched sources", 1)
    for r in report["similarity"]["sources"]:
        where = r.get("doi") and f"doi:{r['doi']}" or r.get("url") or ""
        d.add_paragraph(f"[{r['rank']}] {r['title']} — {r['percent']}% counted, {r['semantic_percent']}% paraphrase. {where}", style="List Number")
    d.add_heading("Findings", 1)
    rank = {r["id"]: r["rank"] for r in report["similarity"]["sources"]}
    for f in report["similarity"]["findings"]:
        p = d.add_paragraph()
        p.add_run(f"[{rank.get(f['source_id'], '?')}] {f['label']} ({f['words']} words, {f['method'].replace('_', ' ')}). ").bold = True
        p.add_run(f["explanation"] + " " + f["next_steps"])
        q = d.add_paragraph(report["text"][f["start"] : f["end"]][:1500])
        q.paragraph_format.left_indent = Pt(18)
        if f.get("source_excerpt"):
            src = d.add_paragraph("Source: " + f["source_excerpt"][:1500])
            src.paragraph_format.left_indent = Pt(18)
            for run in src.runs:
                run.font.color.rgb = RGBColor(0x47, 0x55, 0x69)
    ai = report["ai"]
    d.add_heading("AI-writing indicators", 1)
    d.add_paragraph(f"{ai['label']}. {ai.get('explanation', '')}")
    for c in ai.get("caveats", []):
        d.add_paragraph(c, style="List Bullet")
    d.add_heading("Writing quality", 1)
    d.add_paragraph(f"Quality index {report['quality']['score']}/100 — " + ", ".join(f"{k} {v}" for k, v in report["quality"]["subscores"].items()))
    d.add_heading("Citations and references", 1)
    for issue in report["citations"]["issues"][:60]:
        d.add_paragraph(issue["message"], style="List Bullet")
    d.add_page_break()
    d.add_heading("Document with highlighted matches", 1)
    text = report["text"]
    marks = sorted((f for f in report["similarity"]["findings"]), key=lambda f: f["start"])
    for para_text, offset in _split_with_offsets(text):
        p = d.add_paragraph()
        pos = offset
        end = offset + len(para_text)
        for f in marks:
            if f["end"] <= pos or f["start"] >= end:
                continue
            a, b = max(f["start"], pos), min(f["end"], end)
            if a > pos:
                p.add_run(text[pos:a])
            run = p.add_run(text[a:b])
            run.font.highlight_color = DOCX_HIGHLIGHT.get(f["category"], WD_COLOR_INDEX.GRAY_25)
            pos = b
        if pos < end:
            p.add_run(text[pos:end])
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def _split_with_offsets(text: str) -> list[tuple[str, int]]:
    out, pos = [], 0
    for block in text.split("\n\n"):
        out.append((block, pos))
        pos += len(block) + 2
    return out


def revised_docx(title: str, text: str) -> bytes:
    d = DocxDocument()
    d.styles["Normal"].font.size = Pt(11)
    d.core_properties.title = title
    d.add_heading(title, 1)
    for block in text.split("\n\n"):
        if block.strip():
            d.add_paragraph(block.strip())
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def revised_pdf(title: str, text: str) -> bytes:
    st = _styles()
    buf = io.BytesIO()
    pdf = SimpleDocTemplate(buf, pagesize=A4, leftMargin=22 * mm, rightMargin=22 * mm, topMargin=20 * mm, bottomMargin=20 * mm, title=title)
    story = [_p(title, st["title"]), _p(f"Revised {datetime.utcnow():%Y-%m-%d}", st["small"]), Spacer(1, 8)]
    story += [_p(b.strip(), st["doc"]) for b in text.split("\n\n") if b.strip()]
    pdf.build(story)
    return buf.getvalue()
