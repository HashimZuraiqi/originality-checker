import io
import zipfile

import docx
import pytest
from reportlab.pdfgen import canvas

from app.services.extraction.parsers import ExtractionError
from app.services.extraction.sandbox import extract


def make_docx(paragraphs):
    d = docx.Document()
    for p in paragraphs:
        d.add_paragraph(p)
    t = d.add_table(rows=1, cols=2)
    t.cell(0, 0).text, t.cell(0, 1).text = "Cell A", "Cell B"
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def make_pdf(lines_per_page, pages=1, encrypt=None):
    buf = io.BytesIO()
    c = canvas.Canvas(buf, encrypt=encrypt)
    for pg in range(pages):
        y = 780
        c.drawString(72, 810, "Journal of Testing")
        for line in lines_per_page:
            c.drawString(72, y, line.format(pg=pg))
            y -= 14
        c.drawString(300, 30, str(pg + 1))
        c.showPage()
    c.save()
    return buf.getvalue()


def test_docx_paragraphs_and_tables():
    r = extract("essay.docx", make_docx(["First paragraph here.", "Second paragraph."]))
    assert r.file_type == "docx"
    assert r.text == "First paragraph here.\n\nSecond paragraph.\n\nCell A | Cell B"


def test_txt_encodings():
    assert extract("a.txt", "naïve café".encode()).text == "naïve café"
    assert extract("a.txt", "﻿bom text".encode()).text == "bom text"
    r = extract("a.txt", "caf\xe9 au lait".encode("cp1252"))
    assert "café" in r.text and r.warnings


def test_pdf_text_dehyphenation_and_header_removal():
    lines = ["This is body text on page {pg} with some ordi-", "nary words that continue the sentence nicely.", "Another line of body text follows here, page {pg}."] * 4
    r = extract("paper.pdf", make_pdf(lines, pages=5))
    assert r.pages == 5
    assert "ordinary" in r.text
    assert "Journal of Testing" not in r.text
    assert any("header" in w for w in r.warnings)


def test_encrypted_pdf_rejected():
    data = make_pdf(["Secret text line number {pg} with words."] * 10, encrypt="hunter2")
    with pytest.raises(ExtractionError, match="password"):
        extract("locked.pdf", data)


def test_scanned_pdf_without_text_rejected():
    buf = io.BytesIO()
    c = canvas.Canvas(buf)
    c.rect(50, 50, 200, 200, fill=1)
    c.showPage()
    c.save()
    with pytest.raises(ExtractionError, match="OCR"):
        extract("scan.pdf", buf.getvalue())


@pytest.mark.parametrize(
    "name,data,match",
    [
        ("empty.txt", b"", "empty"),
        ("bad.pdf", b"%PDF-1.7 not really a pdf", "damaged"),
        ("bad.docx", b"PK\x03\x04garbage", "damaged"),
        ("old.doc", b"\xd0\xcf\x11\xe0", ".doc"),
        ("prog.exe", b"MZ\x90\x00", "Unsupported"),
        ("bin.txt", b"\x00\x01\x02binary", "binary"),
        ("evil.txt", b"%PDF-1.4 sneaky", "PDF"),
    ],
)
def test_bad_files_get_clear_errors(name, data, match):
    with pytest.raises(ExtractionError, match=match):
        extract(name, data)


def test_zip_bomb_docx_rejected():
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as zf:
        zf.writestr("word/document.xml", b"<w:document>" + b"A" * (60 * 1024 * 1024) + b"</w:document>")
    with pytest.raises(ExtractionError, match="unsafe size"):
        extract("bomb.docx", buf.getvalue())


def test_size_limit(monkeypatch):
    from app.config import get_settings

    monkeypatch.setattr(get_settings(), "max_upload_mb", 1)
    with pytest.raises(ExtractionError, match="larger"):
        extract("big.txt", b"a " * (600 * 1024))
